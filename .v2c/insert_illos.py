"""Insert the spot illustrations into board_deck.pptx and board_deck.html,
then regenerate board_deck.pdf from the corrected HTML."""
import base64
import os
import subprocess
import sys

from pptx import Presentation
from pptx.util import Inches

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
PRES = os.path.join(ROOT, "presentation")
PNG_DIR = os.path.join(BASE, "png")

EMU = 914400

# name -> (pptx slide index 0-based, left, top, w, h) in inches
PLACEMENTS = {
    "hero_title":     (0, 7.55, 0.98, 5.05, 2.47),
    "ship_leak":      (1, 4.45, 4.55, 3.05, 1.81),
    "gauge_flat":     (3, 8.93, 4.98, 3.70, 1.50),
    "leaky_bucket":   (4, 8.93, 4.98, 3.70, 1.50),
    "globe_pins":     (5, 8.93, 4.98, 3.70, 1.50),
    "table_tag":      (6, 8.93, 4.98, 3.70, 1.50),
    "discount_cliff": (7, 8.93, 4.98, 3.70, 1.50),
    "freight_truck":  (8, 8.93, 4.98, 3.70, 1.50),
    "gavel_check":    (14, 4.45, 3.98, 3.00, 1.62),
    "sail_ship":      (15, 5.07, 1.18, 3.20, 1.56),
}

# html placements: px on the 1280x720 CSS canvas (96 px/inch)
HTML_PLACEMENTS = {
    "hero_title":     (0, 725, 94, 485, 237),
    "ship_leak":      (1, 427, 437, 293, 174),
    "gauge_flat":     (3, 857, 478, 355, 144),
    "leaky_bucket":   (4, 857, 478, 355, 144),
    "globe_pins":     (5, 857, 478, 355, 144),
    "table_tag":      (6, 857, 478, 355, 144),
    "discount_cliff": (7, 857, 478, 355, 144),
    "freight_truck":  (8, 857, 478, 355, 144),
    "gavel_check":    (14, 427, 382, 288, 156),
    "sail_ship":      (15, 487, 113, 307, 150),
}

CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
PROFILE = os.path.join(os.environ.get("TEMP", "/tmp"), "chrome_pdf_profile")


def insert_pptx() -> None:
    path = os.path.join(PRES, "board_deck.pptx")
    prs = Presentation(path)
    by_slide = {}
    for name, (idx, l, t, w, h) in PLACEMENTS.items():
        by_slide.setdefault(idx, []).append(name)
    for idx, names in by_slide.items():
        slide = prs.slides[idx]
        for name in names:
            _, l, t, w, h = PLACEMENTS[name]
            slide.shapes.add_picture(
                os.path.join(PNG_DIR, name + ".png"),
                int(l * EMU), int(t * EMU), int(w * EMU), int(h * EMU),
            )
    prs.save(path)
    print(f"pptx: inserted {len(PLACEMENTS)} illustrations")


def insert_html() -> None:
    path = os.path.join(PRES, "board_deck.html")
    with open(path, encoding="utf-8") as f:
        html = f.read()

    # collect section opening tags in document order
    import re
    opens = list(re.finditer(r"<section class=\"s[^\"]*\"[^>]*>", html))
    assert len(opens) == 16, f"expected 16 sections, found {len(opens)}"

    inserts = []  # (insert_position, img_tag)
    for name, (sec_idx, l, t, w, h) in HTML_PLACEMENTS.items():
        with open(os.path.join(PNG_DIR, name + ".png"), "rb") as f:
            b64 = base64.b64encode(f.read()).decode()
        m = opens[sec_idx]
        tag_end = m.end()
        img = (
            f'<img class="spot" src="data:image/png;base64,{b64}" '
            f'style="position:absolute;left:{l}px;top:{t}px;width:{w}px;height:{h}px;">'
        )
        inserts.append((tag_end, img))

    # insert from the end so earlier offsets stay valid
    for pos, img in sorted(inserts, key=lambda p: -p[0]):
        html = html[:pos] + img + html[pos:]

    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    print("html: inserted 10 data-URI images")


def render_pdf() -> None:
    src = os.path.join(PRES, "board_deck.html")
    tmp_pdf = os.path.join(os.environ.get("TEMP", "/tmp"), "board_deck_new.pdf")
    if os.path.exists(tmp_pdf):
        os.remove(tmp_pdf)
    subprocess.run(
        [CHROME, "--headless=new", "--disable-gpu", "--no-first-run",
         f"--user-data-dir={PROFILE}", "--no-pdf-header-footer",
         f"--print-to-pdf={tmp_pdf}", src],
        capture_output=True, text=True, timeout=180, check=True,
    )
    dst = os.path.join(PRES, "board_deck.pdf")
    size = os.path.getsize(tmp_pdf)
    with open(tmp_pdf, "rb") as fsrc, open(dst, "wb") as fdst:
        fdst.write(fsrc.read())
    print(f"pdf: re-rendered ({size} bytes)")


if __name__ == "__main__":
    insert_pptx()
    insert_html()
    render_pdf()
