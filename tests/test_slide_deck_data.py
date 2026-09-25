"""Unit test suite verifying presentation slide deck data points against dataset."""
import pytest
import os
from src.services.data_loader import load_sales_data
from src.services.analyzer import (
    compute_overview_kpis,
    compute_yoy_growth,
    simulate_turnaround_impact,
)


@pytest.fixture(scope="module")
def global_df():
    """Load the full 51,290 record Global Superstore dataset."""
    return load_sales_data()


def test_slide2_and_3_topline_growth(global_df):
    """Slide 2 & 3: Revenue scaled from $2.26M to $4.30M (+90.3%, CAGR: 23.9%)."""
    yoy = compute_yoy_growth(global_df)
    sales_2011 = yoy.loc[yoy["Year"] == 2011, "Sales"].values[0]
    sales_2014 = yoy.loc[yoy["Year"] == 2014, "Sales"].values[0]
    growth_pct = (sales_2014 - sales_2011) / sales_2011 * 100.0
    cagr = ((sales_2014 / sales_2011) ** (1 / 3) - 1) * 100.0

    assert sales_2011 == pytest.approx(2_259_450.90, rel=1e-3)
    assert sales_2014 == pytest.approx(4_299_865.87, rel=1e-3)
    assert growth_pct == pytest.approx(90.3, abs=0.1)
    assert cagr == pytest.approx(23.9, abs=0.1)


def test_slide2_and_5_profit_leakage(global_df):
    """Slide 2 & 5: $920,646 profit destroyed across 12,544 loss-making lines (24.5%)."""
    kpis = compute_overview_kpis(global_df)
    assert kpis["loss_order_count"] == 12544
    assert kpis["loss_order_pct"] == pytest.approx(24.46, abs=0.1)
    assert kpis["profit_loss_drag"] == pytest.approx(920_646.16, abs=1.0)
    assert kpis["profitable_order_gain"] == pytest.approx(2_388_103.45, abs=1.0)
    assert kpis["total_profit"] == pytest.approx(1_467_457.29, abs=1.0)


def test_slide6_country_deficits(global_df):
    """Slide 6: Turkey (-$98K) and Nigeria (-$81K) lead sovereign deficits."""
    turkey_profit = global_df[global_df["Country"] == "Turkey"]["Profit"].sum()
    nigeria_profit = global_df[global_df["Country"] == "Nigeria"]["Profit"].sum()
    netherlands_profit = global_df[global_df["Country"] == "Netherlands"]["Profit"].sum()
    honduras_profit = global_df[global_df["Country"] == "Honduras"]["Profit"].sum()

    assert turkey_profit == pytest.approx(-98_447.23, abs=1.0)
    assert nigeria_profit == pytest.approx(-80_750.72, abs=1.0)
    assert netherlands_profit == pytest.approx(-41_070.08, abs=1.0)
    assert honduras_profit == pytest.approx(-29_482.37, abs=1.0)
    assert (turkey_profit + nigeria_profit) == pytest.approx(-179_197.95, abs=1.0)


def test_slide7_tables_anomaly(global_df):
    """Slide 7: Tables is the sole net deficit sub-category (-$64.1K on $757K sales)."""
    sub_perf = global_df.groupby("Sub-Category")["Profit"].sum()
    deficit_subs = sub_perf[sub_perf < 0]
    assert len(deficit_subs) == 1
    assert "Tables" in deficit_subs.index
    assert deficit_subs["Tables"] == pytest.approx(-64_083.19, abs=1.0)


def test_slide8_discount_inversion(global_df):
    """Slide 8: Orders > 20% discount destroy $814,682 across 11,328 lines."""
    over_20 = global_df[global_df["Discount"] > 0.20]
    assert len(over_20) == 11328
    assert over_20["Profit"].sum() == pytest.approx(-814_682.09, abs=1.0)


def test_margin_distribution_by_band(global_df):
    """Chapter 3 boxplot evidence: median margin flips negative past 20%; >50% band loses 100%."""
    from src.services.analyzer import analyze_margin_distribution

    dist = {d["band"]: d for d in analyze_margin_distribution(global_df)}
    assert len(dist) == 7
    assert dist["0%"]["median"] == pytest.approx(27.0, abs=1.0)
    assert dist["0%"]["loss_share_pct"] == pytest.approx(0.0, abs=0.5)
    assert dist["10.1-20%"]["median"] == pytest.approx(12.5, abs=1.0)
    assert dist["20.1-30%"]["median"] == pytest.approx(-4.3, abs=1.0)
    assert dist[">50%"]["median"] == pytest.approx(-106.7, abs=2.0)
    assert dist[">50%"]["loss_share_pct"] == pytest.approx(100.0, abs=0.5)
    assert dist[">50%"]["q3"] < 0  # even the top quartile of the >50% band loses money


def test_measure_correlations(global_df):
    """Chapter 3 correlation evidence: Pearson understates the discount cliff (Spearman is stronger)."""
    from src.services.analyzer import analyze_measure_correlations

    corr = analyze_measure_correlations(global_df)
    assert corr["pearson"].loc["Discount", "Profit"] == pytest.approx(-0.316, abs=0.01)
    assert corr["spearman"].loc["Discount", "Profit"] == pytest.approx(-0.596, abs=0.01)
    assert abs(corr["pearson"].loc["Discount", "Sales"]) < 0.15
    assert abs(corr["pearson"].loc["Discount", "Quantity"]) < 0.10
    assert corr["pearson"].loc["Sales", "Profit"] > 0.3  # growth itself is healthy


def test_distribution_correlation_charts(global_df):
    """The two Chapter 3 exhibits construct with the expected traces."""
    from src.components.charts import create_correlation_heatmap_chart, create_margin_boxplot_chart
    from src.services.analyzer import analyze_margin_distribution, analyze_measure_correlations

    fig_box = create_margin_boxplot_chart(analyze_margin_distribution(global_df))
    assert len(fig_box.data) == 7  # one box per discount band

    fig_corr = create_correlation_heatmap_chart(analyze_measure_correlations(global_df))
    assert len(fig_corr.data) == 2  # Pearson + Spearman panels
    assert fig_corr.data[0].type == "heatmap"
    assert len(fig_corr.data[0].z) == 5 and len(fig_corr.data[0].z[0]) == 5


def test_slide12_and_13_base_case_turnaround(global_df):
    """Slide 12 & 13: Base Case turnaround simulation lifts EBITDA +$1.23M to $2.70M (19.9% margin)."""
    sim = simulate_turnaround_impact(
        global_df,
        max_discount_cap=0.20,
        restructure_deficit_territories=True,
        table_freight_surcharge=15.0,
        volume_attrition_rate=0.05,
    )
    assert sim["baseline_profit"] == pytest.approx(1_467_457.29, abs=1.0)
    assert sim["pricing_recovery"] == pytest.approx(1_032_488.44, abs=1.0)
    assert sim["territory_recovery"] == pytest.approx(179_197.95, abs=1.0)
    assert sim["table_freight_recovery"] == pytest.approx(46_245.00, abs=1.0)
    assert sim["attrition_drag"] == pytest.approx(24_127.63, abs=1.0)
    assert sim["net_ebitda_uplift"] == pytest.approx(1_233_803.75, abs=1.0)
    assert sim["projected_profit"] == pytest.approx(2_701_261.04, abs=1.0)
    assert sim["projected_margin"] == pytest.approx(19.89, abs=0.1)


def test_html_slide_no_typo_symbol():
    """Verify HTML slide deck does not contain '$-' currency formatting error."""
    html_path = os.path.join("presentation", "board_deck.html")
    if os.path.exists(html_path):
        with open(html_path, "r", encoding="utf-8") as f:
            content = f.read()
        assert "$-64K" not in content, "Found '$-64K' typo in board_deck.html"
        assert "-$64K" in content, "Expected '-$64K' in board_deck.html"


def test_slide2_operating_profit_visual_and_legend():
    """Verify Slide 2 includes Operating Profit visual points, legend, and volume card."""
    html_path = os.path.join("presentation", "board_deck.html")
    with open(html_path, "r", encoding="utf-8") as f:
        content = f.read()
    assert "$249K" in content
    assert "$504K" in content
    assert "Operating profit (line)" in content
    assert "Order volume (+92.1%)" in content
    assert "Avg order value (flat, -0.9%)" in content


def test_slide7_and_8_legends_and_colors():
    """Verify Slide 7 has discount margin legend and Slide 8 uses var(--rule) matching legend."""
    html_path = os.path.join("presentation", "board_deck.html")
    with open(html_path, "r", encoding="utf-8") as f:
        content = f.read()
    assert "Profitable margin (&le;20% discount)" in content
    assert "Operating deficit (&gt;20% discount)" in content
    # In slide 8, Same Day and First Class should have fill="var(--rule)"
    assert 'fill="var(--rule)"/><text class="sv-val" x="717" y="220" text-anchor="start">17.4%</text>' in content


def test_dashboard_priority_chart_and_discount_cliff_legend(global_df):
    """Verify create_priority_freight_chart and create_discount_cliff_chart have traces with legends."""
    from src.services.analyzer import analyze_discount_impact, analyze_shipping_and_priority
    from src.components.charts import create_discount_cliff_chart, create_priority_freight_chart

    disc_df = analyze_discount_impact(global_df)
    cliff_fig = create_discount_cliff_chart(disc_df)
    assert len(cliff_fig.data) == 2
    assert "Profitable Tier" in cliff_fig.data[0].name
    assert "Deficit Tier" in cliff_fig.data[1].name

    _, priority_df = analyze_shipping_and_priority(global_df)
    priority_fig = create_priority_freight_chart(priority_df)
    assert len(priority_fig.data) == 2 # Bar and line
    assert "Freight Cost Ratio (%)" in priority_fig.data[0].name
    assert "Avg Shipping Cost ($)" in priority_fig.data[1].name


def test_executive_summary_docx():
    """Verify that the 1-page executive summary docx exists and contains all strategic findings and metrics."""
    import os
    docx_path = os.path.join("presentation", "Global_Superstore_Executive_Summary.docx")
    assert os.path.exists(docx_path), "Executive summary .docx file must exist"
    assert os.path.getsize(docx_path) > 10_000, "Executive summary .docx should be non-empty"

    try:
        import docx
        doc = docx.Document(docx_path)
        all_text = " ".join([p.text for p in doc.paragraphs])
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    all_text += " " + cell.text
    except ImportError:
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(docx_path) as z:
            xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            all_text = " ".join(tree.itertext())

    assert "$4.30M" in all_text
    assert "11.6%" in all_text
    assert "-$920,646" in all_text
    assert "$1,233,804" in all_text or "+$1.23M" in all_text
    assert "20% Discount Inversion Cliff" in all_text
    assert "Turkey -$98.4K" in all_text
    assert "Tables Anomaly" in all_text
    assert "Pillar 1: Pricing Governance" in all_text
    assert "Pillar 2: Channel Restructuring" in all_text
    assert "Pillar 3: Freight Surcharges" in all_text
    assert "Resolution 1: ERP Pricing Lock" in all_text
    assert "Resolution 2: Sovereign 3PL Shift" in all_text
    assert "Resolution 3: Governance Charter" in all_text


def test_alternatives_assessment_numbers(global_df):
    """Alternatives-elimination evidence: freight $1.35M (10.7%), deep-discount losses 88.5%, Turkey+Nigeria $163K sales."""
    from src.services.analyzer import compute_alternatives_assessment

    alts = compute_alternatives_assessment(global_df)
    assert alts["freight_total"] == pytest.approx(1_352_820.69, abs=1.0)
    assert alts["freight_ratio"] == pytest.approx(0.107, abs=0.001)
    assert alts["loss_dollars_total"] == pytest.approx(920_646.16, abs=1.0)
    assert alts["deep_discount_loss"] == pytest.approx(814_682.09, abs=1.0)
    assert alts["deep_discount_lines"] == 11328
    assert alts["deep_share_of_losses"] == pytest.approx(0.8849, abs=0.001)
    assert alts["tn_sales"] == pytest.approx(162_858.30, abs=1.0)
    assert alts["tn_loss"] == pytest.approx(179_197.95, abs=1.0)
    assert alts["freight_cut_recovery"] == pytest.approx(135_282.07, abs=1.0)
    assert isinstance(alts["profit_by_band"], list) and len(alts["profit_by_band"]) == 7
    band_net_total = sum(b["net_profit"] for b in alts["profit_by_band"])
    assert band_net_total == pytest.approx(1_467_457.29, abs=1.0)
    above_cap_net = sum(
        b["net_profit"] for b in alts["profit_by_band"] if b["band"] not in ("0%", "0.1-10%", "10.1-20%")
    )
    assert above_cap_net == pytest.approx(-814_682.09, abs=1.0)


def test_dashboard_alternatives_narrative():
    """Verify the revival strategy view renders the alternatives-eliminated section off the analyzer."""
    revival_path = os.path.join("src", "views", "revival_strategy.py")
    with open(revival_path, "r", encoding="utf-8") as f:
        content = f.read()
    assert "Alternatives We Considered and Rejected" in content
    assert "Why Each One Fails" in content
    assert "compute_alternatives_assessment" in content
    assert "Raise list prices, allow deeper discounts" in content
    assert "create_alt1_loss_concentration_chart" in content
    assert "create_alt2_exit_vs_3pl_chart" in content
    assert "create_alt3_freight_cut_vs_leak_chart" in content


def test_alternative_elimination_charts(global_df):
    """Verify the three alternatives-elimination exhibits carry the right traces and evidence."""
    from src.components.charts import (
        create_alt1_loss_concentration_chart,
        create_alt2_exit_vs_3pl_chart,
        create_alt3_freight_cut_vs_leak_chart,
    )
    from src.services.analyzer import compute_alternatives_assessment

    alts = compute_alternatives_assessment(global_df)

    fig1 = create_alt1_loss_concentration_chart(alts["profit_by_band"], alts["deep_share_of_losses"])
    assert len(fig1.data) == 2
    assert "at or below the 20% cap" in fig1.data[0].name
    assert "above the 20% cap" in fig1.data[1].name
    assert sum(fig1.data[0].y) + sum(fig1.data[1].y) == pytest.approx(1_467_457.29, abs=1.0)
    assert sum(fig1.data[1].y) == pytest.approx(-814_682.09, abs=1.0)

    fig2 = create_alt2_exit_vs_3pl_chart(alts["tn_sales"], alts["tn_loss"])
    assert len(fig2.data) == 2
    assert "Exit market entirely" in fig2.data[0].name
    assert "3PL restructuring" in fig2.data[1].name
    assert list(fig2.data[0].y) == [0.0, pytest.approx(179_197.95, abs=1.0)]
    assert list(fig2.data[1].y) == [pytest.approx(162_858.30, abs=1.0), pytest.approx(179_197.95, abs=1.0)]

    fig3 = create_alt3_freight_cut_vs_leak_chart(
        alts["freight_total"], alts["freight_cut_recovery"], alts["deep_discount_loss"]
    )
    assert len(fig3.data) == 1
    assert list(fig3.data[0].y) == [
        pytest.approx(1_352_820.69, abs=1.0),
        pytest.approx(135_282.07, abs=1.0),
        pytest.approx(814_682.09, abs=1.0),
    ]


def test_deck_alternatives_slide():
    """Verify the board deck contains the alternatives-eliminated slide with the tested options."""
    from pptx import Presentation

    prs = Presentation(os.path.join("presentation", "board_deck.pptx"))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text)
    joined = " ".join(texts)
    assert "outearns every alternative tested" in joined
    assert "Raise list prices, allow deeper discounts" in joined
    assert "Exit Turkey and Nigeria entirely" in joined
    assert "Renegotiate carrier rates company-wide" in joined
    assert "Support lever only" in joined


def test_dashboard_narrative_data_accuracy():
    """Verify that all dashboard view narrative files reflect exact empirical calculations without discrepancies."""
    import os

    # 1. Verify eda.py has exact customer segment margins and Tables metrics
    eda_path = os.path.join("src", "views", "eda.py")
    with open(eda_path, "r", encoding="utf-8") as f:
        eda_content = f.read()
    assert "Consumer (11.51%)" in eda_content
    assert "Corporate (11.54%)" in eda_content
    assert "Home Office (11.99%)" in eda_content
    assert "29.1% average promotional discounts" in eda_content
    assert "179,198 in Turkey and Nigeria" in eda_content

    # 2. Verify trends.py clarifies 88.5% loss dollars vs 81.2% transaction volume
    trends_path = os.path.join("src", "views", "trends.py")
    with open(trends_path, "r", encoding="utf-8") as f:
        trends_content = f.read()
    assert "88.5% of all lost dollars" in trends_content
    assert "81.2% of all loss-making lines (10,180)" in trends_content


    # 3. Verify executive_summary.py has exact Base Case turnaround values
    exec_path = os.path.join("src", "views", "executive_summary.py")
    with open(exec_path, "r", encoding="utf-8") as f:
        exec_content = f.read()
    assert "1,233,804 in profit" in exec_content
    assert "1,032,488" in exec_content
    assert "179,198" in exec_content

    # 4. Verify revival_strategy.py has exact Base Case initiative metrics
    revival_path = os.path.join("src", "views", "revival_strategy.py")
    with open(revival_path, "r", encoding="utf-8") as f:
        revival_content = f.read()
    assert "1,032,488 in operating profit" in revival_content
    assert "179,198 in losses" in revival_content
    assert "46,245 from freight fees" in revival_content


def test_deck_presenter_notes():
    """Verify every slide carries presenter notes and the timing budget stays near the 10-minute briefing."""
    import re

    from pptx import Presentation

    prs = Presentation(os.path.join("presentation", "board_deck.pptx"))
    assert len(prs.slides) == 18

    total_seconds = 0.0
    for i, slide in enumerate(prs.slides, 1):
        assert slide.has_notes_slide, f"slide {i} has no notes slide"
        text = slide.notes_slide.notes_text_frame.text
        m = re.match(r"\[(\d+):(\d+) - Slide \d+ of 18\]", text)
        assert m, f"slide {i} notes missing timing header: {text[:40]!r}"
        assert "SAY:" in text, f"slide {i} notes missing talk track"
        total_seconds += int(m.group(1)) * 60 + int(m.group(2))

    assert 9.0 <= total_seconds / 60.0 <= 11.0, f"briefing length {total_seconds / 60.0:.2f} min off budget"


def test_slide_what_happens_next_and_cliff_stats():
    """Deck: the what-happens-next slide exists and the cliff slide carries the new statistical evidence."""
    from pptx import Presentation

    prs = Presentation(os.path.join("presentation", "board_deck.pptx"))
    assert len(prs.slides) == 18

    texts = []
    for slide in prs.slides:
        text = ""
        for shape in slide.shapes:
            if shape.has_text_frame:
                text += " " + shape.text_frame.text
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text += " " + cell.text
        texts.append(text)

    # New slide: forecast, measured cap effect, churn lift
    assert any("What Happens Next" in t and "5.28M" in t for t in texts)
    assert any("101% of volume kept" in t for t in texts)
    assert any("2.9x the average rate" in t for t in texts)
    assert any("AUC 0.80" in t for t in texts)

    # Cliff slide: distribution + correlation evidence
    cliff = next(t for t in texts if "Order economics invert" in t)
    assert "Past 50% off, 100% of orders lose money." in cliff
    assert "-0.60 (Pearson: -0.32)" in cliff

    # Page numbers stayed sequential after the insertion
    numbers = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.name == "TextBox 6" and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text.isdigit():
                    numbers.append(text)
    assert numbers == [str(i) for i in range(1, 17)]



